"""
3회 반복 분류 결과(classification_results_runs3.csv)에서 회차 간 판정이 갈린 논문을 뽑아 사람
리뷰용 엑셀로 정리한다. 최종 라벨은 이미 결과 CSV에 다수결(3회 중 2회 이상 나온 라벨)로
들어 있으므로, 이 스크립트는 그 결과를 바꾸지 않고 "어디서 갈렸는지"만 보여준다.

논문(초록 없는 논문 제외)을 셋으로 나눈다.
- 완전일치: 3회 판정(코드 집합 또는 ERR)이 모두 같음 -> 그대로 채택
- 다수결: 3회 판정이 일부 달랐지만 2회 이상 나온 라벨이 있음 -> 다수결 라벨 채택
- 과반없음: status는 OK(ERR이 1회 이하)인데 2회 이상 나온 라벨이 하나도 없음 -> 라벨 빈칸,
  사람이 최종 판정해야 함

결과 CSV의 rationale 칸은 3회분 근거를 " | "로 이어 붙인 것인데, 한 회차 안에서도 ROLE 근거와
ENV 근거가 " | "로 붙어 있다. 그래서 "[ENV]"로 시작하는 조각이 바로 앞 "[ROLE]" 조각에 이어지면
같은 회차로 묶는다(_split_runs). 현재 결과 685편 전부 정확히 3회분으로 나뉘는 것을 확인했다.

사용법:
    python3 target_new/analyze_runs3.py
    (target_new/리뷰_3회불일치.xlsx 생성 - 시트: 요약 / 과반없음 / 다수결 / 전체)
    (target_new/리뷰대상_과반없음_N편.xlsx 생성 - 과반없음 논문만 담은 리뷰 전용 파일)
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd

HERE = Path(__file__).parent
RUNS3_CSV = HERE / "classification_results_runs3.csv"
RUN1_CSV = HERE / "classification_results.csv"
XLSX_PATH = HERE / "KCI 사회과학.xlsx"
OUT_XLSX = HERE / "리뷰_3회불일치.xlsx"
REVIEW_XLSX = HERE / "리뷰대상_과반없음.xlsx"  # 실제 파일명에는 편수가 붙는다
CODES = [
    "ROLE_FIELD_PUBLICSPHERE", "ROLE_FIELD_SOCIALCAPITAL",
    "ROLE_WINDOW_POLARIZE", "ROLE_WINDOW_GENDER", "ROLE_WINDOW_HATE", "ROLE_WINDOW_MISC",
    "ENV_COMMUNITY_SUBSCRIPTION", "ENV_COMMUNITY_GOVERNANCE", "ENV_COMMUNITY_DEMOGRAPHIC",
    "ENV_COMMUNITY_LIFECYCLE", "ENV_ONLINE_ANNONIMITY",
]


def _split_runs(text: str) -> list[str]:
    runs: list[str] = []
    for token in str(text).split(" | "):
        if token.startswith("[ENV]") and runs and runs[-1].startswith("[ROLE]") and "[ENV]" not in runs[-1]:
            runs[-1] += " | " + token
        else:
            runs.append(token)
    return runs


def _has_label(row: pd.Series) -> bool:
    return any(str(row[c]).strip() not in ("", "nan") for c in ("role_codes", "env_codes"))


def _agreement(row: pd.Series, run_codes: list[str]) -> str:
    if len(set(run_codes)) == 1:
        return "완전일치"
    if row["status"] == "OK" and not _has_label(row):
        return "과반없음"
    return "다수결"


def _write_review_file(no_majority: pd.DataFrame, path: Path) -> None:
    """과반없음 논문만 담은 리뷰 전용 엑셀. 첫 시트(리뷰)에 한 논문씩 초록과 회차별 판정·근거를
    나란히 두고, 리뷰어가 채울 칸(최종판정 1·2, 메모)을 맨 앞쪽에 둔다. 최종판정 칸은 코드 목록
    드롭다운이지만 목록 밖 값도 입력할 수 있다. 두 번째 시트(안내)에 입력 규칙과 코드 설명을 둔다."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    cols = ["id", "분야", "title", "published_year", "리뷰어_최종판정1", "리뷰어_최종판정2", "메모",
            "초록", "run1_판정", "run2_판정", "run3_판정", "run1_근거", "run2_근거", "run3_근거", "기존1회_판정"]
    sheet = no_majority.assign(리뷰어_최종판정1="", 리뷰어_최종판정2="", 메모="")[cols]
    sheet = sheet.rename(columns={"title": "제목", "published_year": "연도"})

    guide = pd.DataFrame(
        {
            "항목": ["무엇을 하나요", "최종판정1", "최종판정2", "ERR", "메모", "참고 칸"],
            "설명": [
                "3회 분류가 모두 다르게 나와 다수결로 정하지 못한 논문입니다. 초록과 회차별 근거를 보고 최종 라벨을 정해 주세요.",
                "가장 알맞은 코드 1개. 드롭다운에서 고르거나 직접 입력합니다.",
                "두 번째 코드가 있으면 입력(최대 2개). 없으면 비워 둡니다.",
                "온라인 커뮤니티 연구가 아니라고 판단하면 최종판정1에 ERR을 입력합니다.",
                "판단 근거나 애매한 점을 자유롭게 적습니다.",
                "run1~3_판정/근거는 모델의 3회 판정, 기존1회_판정은 처음 1회 분류 결과입니다.",
            ],
        }
    )

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        sheet.to_excel(writer, sheet_name="리뷰", index=False)
        guide.to_excel(writer, sheet_name="안내", index=False)
        ws = writer.sheets["리뷰"]

        widths = {"id": 14, "분야": 10, "제목": 36, "연도": 7, "리뷰어_최종판정1": 28, "리뷰어_최종판정2": 28,
                  "메모": 30, "초록": 70, "run1_판정": 26, "run2_판정": 26, "run3_판정": 26,
                  "run1_근거": 60, "run2_근거": 60, "run3_근거": 60, "기존1회_판정": 26}
        header_fill = PatternFill("solid", fgColor="08306B")
        input_fill = PatternFill("solid", fgColor="FFF2CC")  # 리뷰어가 채울 칸은 연노랑
        for i, name in enumerate(sheet.columns, 1):
            letter = get_column_letter(i)
            ws.column_dimensions[letter].width = widths.get(name, 15)
            ws[f"{letter}1"].font = Font(bold=True, color="FFFFFF")
            ws[f"{letter}1"].fill = header_fill
            for row in range(2, len(sheet) + 2):
                cell = ws[f"{letter}{row}"]
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if name.startswith("리뷰어_") or name == "메모":
                    cell.fill = input_fill
        for row in range(2, len(sheet) + 2):
            ws.row_dimensions[row].height = 180
        ws.freeze_panes = "E2"  # id~연도는 고정해서 옆으로 스크롤해도 어떤 논문인지 보이게

        dv = DataValidation(type="list", formula1='"' + ",".join(["ERR", *CODES]) + '"', allow_blank=True)
        dv.showErrorMessage = False  # 목록 밖 값(메모성 입력 등)도 막지 않는다
        ws.add_data_validation(dv)
        for letter in ("E", "F"):
            dv.add(f"{letter}2:{letter}{len(sheet) + 1}")

        wg = writer.sheets["안내"]
        wg.column_dimensions["A"].width = 16
        wg.column_dimensions["B"].width = 100
        for row in wg.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")


def main() -> None:
    runs3 = pd.read_csv(RUNS3_CSV, encoding="utf-8-sig", dtype=str)
    runs3 = runs3[runs3["status"] != "NO_ABSTRACT"].copy()

    run1 = pd.read_csv(RUN1_CSV, encoding="utf-8-sig", dtype=str)
    run1["기존1회_판정"] = run1.apply(
        lambda r: "ERR" if r["status"] == "ERR" else ", ".join(
            c for c in (str(r["role_codes"]), str(r["env_codes"])) if c not in ("", "nan")
        ),
        axis=1,
    )
    meta = pd.read_excel(XLSX_PATH, sheet_name="Sheet1", dtype={"논문ID": str})
    meta = meta.rename(columns={"논문ID": "id", "중분류": "분야"})
    meta["초록"] = meta["KOR_ABST"].where(meta["KOR_ABST"].notna(), meta["ENG_ABST"])

    rows = []
    for _, r in runs3.iterrows():
        run_codes = [part.split(":", 1)[1] for part in str(r["run_details"]).split("; ")]
        run_rationales = _split_runs(r["rationale"])
        final = "ERR" if r["status"] == "ERR" else ", ".join(
            c for c in (str(r["role_codes"]), str(r["env_codes"])) if c not in ("", "nan")
        )
        row = {
            "id": r["id"],
            "title": r["title"],
            "published_year": r["published_year"],
            "일치도": _agreement(r, run_codes),
            "최종(다수결)": final or "(빈칸 - 리뷰 필요)",
        }
        for i, (codes, rationale) in enumerate(zip(run_codes, run_rationales), 1):
            row[f"run{i}_판정"] = "ERR/없음" if codes == "없음" else codes
            row[f"run{i}_근거"] = rationale
        rows.append(row)

    out = pd.DataFrame(rows)
    out = out.merge(meta[["id", "분야", "초록"]], on="id", how="left")
    out = out.merge(run1[["id", "기존1회_판정"]], on="id", how="left")
    lead = ["id", "분야", "title", "published_year", "일치도", "최종(다수결)", "기존1회_판정"]
    out = out[lead + [c for c in out.columns if c not in lead]]

    summary = out["일치도"].value_counts().reindex(["완전일치", "다수결", "과반없음"]).rename("편수").to_frame()
    summary["비율(%)"] = (summary["편수"] / summary["편수"].sum() * 100).round(1)

    no_majority = out[out["일치도"] == "과반없음"].copy()
    no_majority["리뷰어_최종판정"] = ""
    no_majority["메모"] = ""

    with pd.ExcelWriter(OUT_XLSX) as writer:
        summary.to_excel(writer, sheet_name="요약")
        no_majority.to_excel(writer, sheet_name="과반없음", index=False)
        out[out["일치도"] == "다수결"].to_excel(writer, sheet_name="다수결", index=False)
        out.to_excel(writer, sheet_name="전체", index=False)

    print(summary.to_string())
    print(f"\n과반없음 {len(no_majority)}편 - 분야별:")
    print(no_majority["분야"].value_counts().to_string())
    review_path = REVIEW_XLSX.with_name(f"리뷰대상_과반없음_{len(no_majority)}편.xlsx")
    _write_review_file(no_majority, review_path)
    print(f"\n저장: {OUT_XLSX}")
    print(f"리뷰 전용: {review_path}")


if __name__ == "__main__":
    main()
